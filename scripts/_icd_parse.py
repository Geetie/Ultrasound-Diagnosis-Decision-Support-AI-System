import openpyxl

wb = openpyxl.load_workbook(
    "data/medical_resources/icd10/国家临床版2.0疾病诊断编码（ICD-10）.xlsx",
    read_only=True,
)
ws = wb.active
print("sheet:", ws.title)
rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
for r in rows:
    print(r)
n = sum(1 for _ in ws.iter_rows(min_row=2, values_only=True))
print("data rows:", n)
